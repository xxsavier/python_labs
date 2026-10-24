from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str | Path, xlsx_path: str | Path) -> None:
    """Конвертирует CSV в XLSX (openpyxl).

    Первая строка CSV это заголовок. Лист называется "Sheet1".
    Ширина колонки равна длине самого длинного текста в ней, но не меньше 8.

    Raises:
        ValueError: неверный тип файла или пустой CSV.
        FileNotFoundError: CSV-файл не найден.
    """
    if Path(csv_path).suffix.lower() != ".csv":
        raise ValueError("Входной файл должен иметь расширение .csv")
    if Path(xlsx_path).suffix.lower() != ".xlsx":
        raise ValueError("Выходной файл должен иметь расширение .xlsx")

    with Path(csv_path).open(encoding="utf-8", newline="") as f:
        rows = [row for row in csv.reader(f) if row]
    if not rows:
        raise ValueError("CSV пустой")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)

    n_cols = max(len(row) for row in rows)
    for col in range(n_cols):
        longest = max((len(row[col]) for row in rows if col < len(row)), default=0)
        ws.column_dimensions[get_column_letter(col + 1)].width = max(8, longest)

    Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)